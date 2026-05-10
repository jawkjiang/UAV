"""Markdown -> .docx converter for the response letter.

Reads ``docs/response_to_reviewers.md`` and emits ``docs/response_to_reviewers.docx``.
Supports:
  - Headings (#, ##, ###) -> Times New Roman bold at 16/13/11 pt
  - Bold (**...**), italic (*...*), bold-italic (***...***), code (`...`) inline runs
  - Blockquotes (lines beginning with >) -> indented italic
  - GitHub-flavoured tables (| a | b | with |---|---| separator)
      -> plain Word Table Grid: bold header row, single-line borders, no banded fills, no color
  - Image embeds ![alt](path) -> centred picture scaled to 5.5 inch width
      (PDF images resolved via sibling .png; PyMuPDF used as fallback)
  - Fenced code blocks -> Consolas 9pt
  - Bullet and numbered lists
  - [Reviewer N, Comment M] headings rendered in red
  - **Action:** paragraphs rendered in blue
  - Everything else black Times New Roman 11pt
"""

import os
import re
import sys
from typing import List, Optional, Tuple

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
SRC = os.path.join(HERE, "response_to_reviewers.md")
DST = os.path.join(HERE, "response_to_reviewers.docx")

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
REV_COLOR = RGBColor(0xB4, 0x3C, 0x3C)
ACT_COLOR = RGBColor(0x00, 0x64, 0xB4)
BLACK = RGBColor(0x00, 0x00, 0x00)

INLINE_PATTERN = re.compile(
    r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)"
)
IMAGE_PATTERN = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.*)$")
TABLE_SEP_PATTERN = re.compile(r"^\|?\s*[:\- ]+\s*(\|\s*[:\- ]+\s*)+\|?\s*$")


def _set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in (qn("w:ascii"), qn("w:hAnsi"), qn("w:cs")):
        rfonts.set(attr, BODY_FONT)


def _frun(paragraph, text: str, bold=False, italic=False,
          size_pt=11, color=None, font=None) -> None:
    r = paragraph.add_run(text)
    r.font.name = font or BODY_FONT
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color if color is not None else BLACK


def _add_runs(paragraph, text: str, base_bold=False, base_italic=False,
              base_color=None) -> None:
    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            _frun(paragraph, text[pos:match.start()],
                  bold=base_bold, italic=base_italic, color=base_color)
        token = match.group(1)
        if token.startswith("***"):
            _frun(paragraph, token[3:-3], bold=True, italic=True, color=base_color)
        elif token.startswith("**"):
            _frun(paragraph, token[2:-2], bold=True, italic=base_italic,
                  color=base_color)
        elif token.startswith("*"):
            _frun(paragraph, token[1:-1], bold=base_bold, italic=True,
                  color=base_color)
        elif token.startswith("`"):
            _frun(paragraph, token[1:-1], font=CODE_FONT, size_pt=10, color=BLACK)
        pos = match.end()
    if pos < len(text):
        _frun(paragraph, text[pos:], bold=base_bold, italic=base_italic,
              color=base_color)


def _remove_cell_shading(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for shd in tcPr.findall(qn("w:shd")):
                tcPr.remove(shd)


def _add_table(doc: Document, header: List[str], rows: List[List[str]]) -> None:
    if not header:
        return
    ncols = len(header)
    table = doc.add_table(rows=1 + len(rows), cols=ncols)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for j, h in enumerate(header):
        p = hdr_cells[j].paragraphs[0]
        _frun(p, h.strip(), bold=True)
    for r_idx, row in enumerate(rows, start=1):
        cells = table.rows[r_idx].cells
        for c_idx in range(ncols):
            p = cells[c_idx].paragraphs[0]
            cell_text = row[c_idx] if c_idx < len(row) else ""
            _add_runs(p, cell_text.strip())
    _remove_cell_shading(table)


def _split_table_row(line: str) -> List[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _is_table_separator(line: str) -> bool:
    return bool(TABLE_SEP_PATTERN.match(line.strip()))


def _pdf_to_png(pdf_path: str) -> Optional[str]:
    sibling = pdf_path[:-4] + ".png"
    if os.path.exists(sibling):
        return sibling
    try:
        import fitz
    except ImportError:
        return None
    try:
        doc = fitz.open(pdf_path)
        page = doc.load_page(0)
        pix = page.get_pixmap(dpi=180)
        pix.save(sibling)
        doc.close()
        return sibling
    except Exception as exc:
        print(f"[warn] PDF->PNG failed for {pdf_path}: {exc}", file=sys.stderr)
        return None


def _resolve_image(rel_path: str) -> Optional[str]:
    for base in (ROOT, HERE):
        candidate = os.path.normpath(os.path.join(base, rel_path))
        if os.path.exists(candidate):
            if candidate.lower().endswith(".pdf"):
                return _pdf_to_png(candidate)
            return candidate
    return None


def convert(md_path: str = SRC, docx_path: str = DST) -> Tuple[int, int]:
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    _set_normal_style(doc)
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    n_tables = 0
    n_images = 0
    i = 0
    in_code = False
    code_buf: List[str] = []

    while i < len(lines):
        line = lines[i]

        # Fenced code block
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                r = p.add_run("\n".join(code_buf))
                r.font.name = CODE_FONT
                r.font.size = Pt(9)
                r.font.color.rgb = BLACK
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Image embed
        m_img = IMAGE_PATTERN.match(line.strip())
        if m_img:
            alt, path = m_img.group(1), m_img.group(2)
            resolved = _resolve_image(path)
            if resolved is not None:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                try:
                    run.add_picture(resolved, width=Inches(5.5))
                    n_images += 1
                except Exception as exc:
                    print(f"[warn] could not embed {resolved}: {exc}", file=sys.stderr)
                if alt:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _frun(cap, alt, italic=True, size_pt=9)
            else:
                print(f"[warn] missing image {path}", file=sys.stderr)
                p = doc.add_paragraph()
                _frun(p, f"[Figure: {alt} — {path} not found]", italic=True)
            i += 1
            continue

        # Heading
        m_h = HEADING_PATTERN.match(line)
        if m_h:
            level = len(m_h.group(1))
            heading_text = m_h.group(2).strip()
            p = doc.add_paragraph()
            size = {1: 16, 2: 13, 3: 11}.get(level, 11)
            color = REV_COLOR if heading_text.startswith("[Reviewer") else None
            _frun(p, heading_text, bold=True, size_pt=size, color=color)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Horizontal rule
        if line.strip() == "---":
            doc.add_paragraph()
            i += 1
            continue

        # Table: header row, separator row, body rows
        if line.strip().startswith("|") and i + 1 < len(lines) \
                and _is_table_separator(lines[i + 1]):
            header = _split_table_row(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            _add_table(doc, header, rows)
            n_tables += 1
            continue

        # Blockquote
        if line.strip().startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].lstrip("> ").rstrip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            _add_runs(p, " ".join(quote_lines), base_italic=True)
            continue

        # Bullet list
        if re.match(r"^\s*[-*]\s+", line):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                content = re.sub(r"^\s*[-*]\s+", "", lines[i])
                p = doc.add_paragraph(style="List Bullet")
                _add_runs(p, content)
                i += 1
            continue

        # Numbered list
        if re.match(r"^\s*\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                content = re.sub(r"^\s*\d+\.\s+", "", lines[i])
                p = doc.add_paragraph(style="List Number")
                _add_runs(p, content)
                i += 1
            continue

        # Blank line
        if line.strip() == "":
            i += 1
            continue

        # Paragraph
        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() != "" \
                and not HEADING_PATTERN.match(lines[i]) \
                and not lines[i].strip().startswith(">") \
                and not lines[i].strip().startswith("|") \
                and not re.match(r"^\s*[-*]\s+", lines[i]) \
                and not re.match(r"^\s*\d+\.\s+", lines[i]) \
                and not lines[i].strip().startswith("```") \
                and not IMAGE_PATTERN.match(lines[i].strip()):
            para_lines.append(lines[i])
            i += 1
        text = " ".join(s.strip() for s in para_lines).strip()
        if text:
            p = doc.add_paragraph()
            stripped = text.lstrip()
            base_color = ACT_COLOR if stripped.startswith("**Action:**") else None
            _add_runs(p, text, base_color=base_color)

    doc.save(docx_path)
    return n_tables, n_images


if __name__ == "__main__":
    n_tables, n_images = convert()
    size = os.path.getsize(DST)
    print(f"Wrote {DST}: {size} bytes, {n_tables} tables, {n_images} images")
